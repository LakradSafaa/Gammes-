from __future__ import annotations

import base64
import os
from io import BytesIO
from typing import Iterable

import qrcode
from django.conf import settings
from django.db import connection
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


STATUT_LABELS = {
    "brouillon": "Brouillon",
    "en_validation": "En validation",
    "validee": "Validée",
    "archivee": "Archivée",
}

TYPE_MAINTENANCE_LABELS = {
    "preventif": "Préventive",
    "correctif": "Corrective",
    "amelioratif": "Améliorative",
    "conditionnel": "Conditionnelle",
    "predictif": "Prédictive",
}

TYPE_ARRET_LABELS = {
    "aucun": "Aucun arrêt",
    "equipement": "Arrêt de l'équipement",
    "partiel": "Arrêt partiel",
    "ligne": "Arrêt de ligne",
    "total": "Arrêt total",
}


# -----------------------------------------------------------------------------
# Helpers
# -----------------------------------------------------------------------------

def _text(value, default: str = "-") -> str:
    if value is None:
        return default
    value = str(value).strip()
    return value if value else default


def _manager_items(obj, manager_name: str):
    """Return a list from a related manager without making export fragile."""
    manager = getattr(obj, manager_name, None)
    if manager is None:
        return []
    try:
        return list(manager.all())
    except Exception:
        return []


def _ordered_manager_items(obj, manager_name: str, *fields: str):
    manager = getattr(obj, manager_name, None)
    if manager is None:
        return []
    try:
        queryset = manager.all()
        if fields:
            queryset = queryset.order_by(*fields)
        return list(queryset)
    except Exception:
        return _manager_items(obj, manager_name)


def _association_label(association, attribute: str, fallback: str = "-") -> str:
    value = getattr(association, attribute, None)
    if value is None:
        return fallback
    return _text(getattr(value, "nom", value), fallback)


def _image_bytes_from_value(value: str | None) -> bytes | None:
    """Read data URLs and local media paths. Remote HTTP files are not fetched."""
    if not value:
        return None

    raw = str(value).strip()
    if not raw:
        return None

    if raw.startswith("data:image/") and "," in raw:
        try:
            encoded = raw.split(",", 1)[1]
            return base64.b64decode(encoded)
        except Exception:
            return None

    candidates = []
    if os.path.isabs(raw):
        candidates.append(raw)
    else:
        media_root = str(getattr(settings, "MEDIA_ROOT", "") or "")
        base_dir = str(getattr(settings, "BASE_DIR", "") or "")
        if media_root:
            candidates.append(os.path.join(media_root, raw.lstrip("/\\")))
        if base_dir:
            candidates.append(os.path.join(base_dir, raw.lstrip("/\\")))

    for candidate in candidates:
        try:
            if os.path.isfile(candidate):
                with open(candidate, "rb") as handle:
                    return handle.read()
        except OSError:
            continue

    return None


def get_statut_label(version) -> str:
    statut = getattr(version, "statut", None)
    if not statut:
        return "-"
    return STATUT_LABELS.get(
        statut,
        str(statut).replace("_", " ").capitalize(),
    )


def get_type_maintenance_label(version) -> str:
    value = getattr(version, "type_maintenance", None)
    if not value:
        return "-"
    return TYPE_MAINTENANCE_LABELS.get(
        value,
        str(value).replace("_", " ").capitalize(),
    )


def get_type_arret_label(version) -> str:
    value = getattr(version, "type_arret", None)
    if not value:
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    "SELECT type_arret FROM gamme_versions WHERE id = %s",
                    [version.id],
                )
                row = cursor.fetchone()
                value = row[0] if row else None
        except Exception:
            value = None
    if value:
        return TYPE_ARRET_LABELS.get(
            value,
            str(value).replace("_", " ").capitalize(),
        )
    return "Arrêt requis" if bool(getattr(version, "arret", False)) else "Aucun arrêt"


def qr_png_bytes(version) -> bytes:
    app_base_url = getattr(
        settings,
        "APP_BASE_URL",
        "http://localhost:5173",
    ).rstrip("/")

    url = f"{app_base_url}/qr/{version.gamme.code}"
    image = qrcode.make(url)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _equipment(version):
    gamme = getattr(version, "gamme", None)
    return getattr(gamme, "equipement", None) if gamme is not None else None


def _list_names(version, manager_name: str, child_name: str) -> list[str]:
    result: list[str] = []
    items = _manager_items(version, manager_name)
    for association in items:
        result.append(_association_label(association, child_name))

    if not result and manager_name == "version_epcs":
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    SELECT e.nom
                    FROM version_epcs ve
                    JOIN epcs e ON e.id = ve.epc_id
                    WHERE ve.version_id = %s
                    ORDER BY e.nom
                    """,
                    [version.id],
                )
                result = [row[0] for row in cursor.fetchall()]
        except Exception:
            result = []
    return result


def _list_quantified(version, manager_name: str, child_name: str) -> list[str]:
    result: list[str] = []
    for association in _manager_items(version, manager_name):
        name = _association_label(association, child_name)
        quantity = getattr(association, "quantite", None)
        result.append(f"{name} × {quantity if quantity is not None else 1}")
    return result


def _filename(version, extension: str) -> str:
    return f"{version.gamme.code}_{version.code_version}.{extension}"


# -----------------------------------------------------------------------------
# PDF - ReportLab, no native GTK/Pango dependency, no local persistent storage
# -----------------------------------------------------------------------------

def generate_pdf(version) -> tuple[bytes, str]:
    """Generate a PDF entirely in memory and return (bytes, filename)."""
    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=13 * mm,
        leftMargin=13 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
        title=f"Gamme {version.gamme.code} {version.code_version}",
        author="Gammes Maintenance",
    )

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="GMTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#063D32"),
            alignment=TA_CENTER,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="GMHeading",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#063D32"),
            spaceBefore=8,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="GMBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#172B2A"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="GMSmall",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9.5,
            textColor=colors.HexColor("#50665F"),
        )
    )

    story = []
    equipment = _equipment(version)

    # Header with QR on first page.
    try:
        qr_image = Image(BytesIO(qr_png_bytes(version)), width=24 * mm, height=24 * mm)
    except Exception:
        qr_image = Paragraph("QR indisponible", styles["GMSmall"])

    title = Paragraph(
        f"<b>Gamme opératoire</b><br/>{_text(version.gamme.code)} — {_text(getattr(version.gamme, 'designation', None))}<br/>"
        f"<font size='9'>{_text(version.code_version)} · {get_statut_label(version)}</font>",
        styles["GMTitle"],
    )

    header = Table([[title, qr_image]], colWidths=[153 * mm, 28 * mm])
    header.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.8, colors.HexColor("#DDE7E3")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.extend([header, Spacer(1, 6)])

    info_rows = [
        ["Code", _text(version.gamme.code), "Version", _text(version.code_version)],
        ["Abréviation", _text(getattr(version.gamme, "abreviation", None)), "Date", _text(getattr(version, "date_version", None))],
        ["Équipement", _text(getattr(equipment, "nom", None)), "Code équipement", _text(getattr(equipment, "code", None))],
        ["Constructeur", _text(getattr(equipment, "constructeur", None)), "Référence", _text(getattr(equipment, "reference", None))],
        ["Type machine", _text(getattr(equipment, "type", None)), "Durée", f"{getattr(version, 'duree_minutes', 0) or 0} min"],
        ["Maintenance", get_type_maintenance_label(version), "Périodicité", _text(getattr(version, "periodicite", None))],
        ["Main-d'œuvre", _text(getattr(version, "main_oeuvre", None), "0"), "Type d'arrêt", get_type_arret_label(version)],
        ["Rédacteur", _text(getattr(version, "redacteur", None)), "Valideur", _text(getattr(version, "valideur", None))],
    ]

    info = Table(info_rows, colWidths=[29 * mm, 62 * mm, 29 * mm, 62 * mm])
    info.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.8),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1FBF7")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F1FBF7")),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#063D32")),
                ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor("#063D32")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#DDE7E3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(info)

    description = getattr(version.gamme, "description", None)
    if description:
        story.extend(
            [
                Paragraph("Description", styles["GMHeading"]),
                Paragraph(_text(description), styles["GMBody"]),
            ]
        )

    modifications = getattr(version, "modifications", None)
    if modifications:
        story.extend(
            [
                Paragraph("Objet / modifications", styles["GMHeading"]),
                Paragraph(_text(modifications), styles["GMBody"]),
            ]
        )

    def add_list_section(title_text: str, items: Iterable[str], empty_text: str):
        story.append(Paragraph(title_text, styles["GMHeading"]))
        values = [value for value in items if value]
        if not values:
            story.append(Paragraph(empty_text, styles["GMSmall"]))
            return
        for value in values:
            story.append(Paragraph(f"• {value}", styles["GMBody"]))

    add_list_section(
        "Équipements de protection individuelle (EPI)",
        _list_names(version, "version_epis", "epi"),
        "Aucun EPI renseigné.",
    )
    add_list_section(
        "Équipements de protection collective (EPC)",
        _list_names(version, "version_epcs", "epc"),
        "Aucun EPC renseigné.",
    )
    add_list_section(
        "Risques",
        _list_names(version, "version_risques", "risque"),
        "Aucun risque renseigné.",
    )
    add_list_section(
        "Outillages",
        _list_quantified(version, "version_outillages", "outillage"),
        "Aucun outillage renseigné.",
    )
    add_list_section(
        "Pièces de rechange",
        _list_quantified(version, "version_pieces_rechange", "piece"),
        "Aucune pièce de rechange renseignée.",
    )

    story.append(Paragraph("Étapes & actions", styles["GMHeading"]))
    etapes = _ordered_manager_items(version, "etapes", "ordre", "numero")
    if not etapes:
        story.append(Paragraph("Aucune étape renseignée.", styles["GMSmall"]))
    else:
        for etape in etapes:
            story.append(
                Paragraph(
                    f"<b>{getattr(etape, 'numero', '-')}. {_text(getattr(etape, 'titre', None))}</b> — "
                    f"{getattr(etape, 'duree_minutes', 0) or 0} min",
                    styles["GMBody"],
                )
            )
            if getattr(etape, "description", None):
                story.append(Paragraph(_text(etape.description), styles["GMSmall"]))

            actions = _ordered_manager_items(etape, "actions", "ordre")
            for action in actions:
                story.append(
                    Paragraph(
                        f"&nbsp;&nbsp;• {_text(getattr(action, 'contenu', None))}",
                        styles["GMBody"],
                    )
                )

            images = _ordered_manager_items(etape, "images", "ordre")
            for image in images:
                image_value = getattr(image, "image_url", None)
                raw_image = _image_bytes_from_value(image_value)
                if raw_image:
                    try:
                        img = Image(BytesIO(raw_image))
                        img._restrictSize(155 * mm, 75 * mm)
                        story.append(Spacer(1, 3))
                        story.append(img)
                    except Exception:
                        pass
                if getattr(image, "description", None):
                    story.append(Paragraph(_text(image.description), styles["GMSmall"]))
            story.append(Spacer(1, 5))

    recommendations = _manager_items(version, "recommandations")
    story.append(Paragraph("Recommandations", styles["GMHeading"]))
    if recommendations:
        for recommendation in recommendations:
            title_text = _text(getattr(recommendation, "titre", None), "Recommandation")
            content = _text(getattr(recommendation, "contenu", None), "")
            story.append(Paragraph(f"<b>{title_text}</b> — {content}", styles["GMBody"]))
    else:
        story.append(Paragraph("Aucune recommandation renseignée.", styles["GMSmall"]))

    documents = _manager_items(version, "documents")
    story.append(Paragraph("Documents liés", styles["GMHeading"]))
    if documents:
        for linked_document in documents:
            text = _text(getattr(linked_document, "titre", None))
            reference = getattr(linked_document, "reference", None)
            if reference:
                text += f" — Réf. {reference}"
            story.append(Paragraph(f"• {text}", styles["GMBody"]))
    else:
        story.append(Paragraph("Aucun document lié.", styles["GMSmall"]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#DDE7E3"))
        canvas.line(13 * mm, 11 * mm, 197 * mm, 11 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#667A74"))
        canvas.drawString(13 * mm, 7 * mm, f"{version.gamme.code} — {version.code_version}")
        canvas.drawRightString(197 * mm, 7 * mm, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue(), _filename(version, "pdf")


# -----------------------------------------------------------------------------
# WORD - in-memory generation, no local persistent storage
# -----------------------------------------------------------------------------

def generate_docx(version) -> tuple[bytes, str]:
    document = Document()
    equipment = _equipment(version)

    title = document.add_heading(
        f"Gamme opératoire — {version.gamme.code}",
        level=0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        f"{_text(getattr(version.gamme, 'designation', None))}\n"
        f"{_text(version.code_version)} · {get_statut_label(version)}"
    )

    # QR immediately after title, so it stays on the first page.
    try:
        document.add_picture(BytesIO(qr_png_bytes(version)), width=Inches(1.25))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    except Exception:
        pass

    document.add_heading("Informations générales", level=1)
    table = document.add_table(rows=0, cols=4)
    table.style = "Table Grid"

    rows = [
        ("Code", version.gamme.code, "Version", version.code_version),
        ("Abréviation", getattr(version.gamme, "abreviation", None), "Date", getattr(version, "date_version", None)),
        ("Équipement", getattr(equipment, "nom", None), "Code équipement", getattr(equipment, "code", None)),
        ("Constructeur", getattr(equipment, "constructeur", None), "Référence", getattr(equipment, "reference", None)),
        ("Type machine", getattr(equipment, "type", None), "Durée", f"{getattr(version, 'duree_minutes', 0) or 0} min"),
        ("Maintenance", get_type_maintenance_label(version), "Périodicité", getattr(version, "periodicite", None)),
        ("Main-d'œuvre", getattr(version, "main_oeuvre", 0), "Type d'arrêt", get_type_arret_label(version)),
        ("Rédacteur", getattr(version, "redacteur", None), "Valideur", getattr(version, "valideur", None)),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = _text(value)
        cells[0].paragraphs[0].runs[0].bold = True
        cells[2].paragraphs[0].runs[0].bold = True

    if getattr(version.gamme, "description", None):
        document.add_heading("Description", level=1)
        document.add_paragraph(_text(version.gamme.description))

    if getattr(version, "modifications", None):
        document.add_heading("Objet / modifications", level=1)
        document.add_paragraph(_text(version.modifications))

    def docx_list(title_text: str, values: Iterable[str], empty_text: str):
        document.add_heading(title_text, level=1)
        values_list = [value for value in values if value]
        if values_list:
            for value in values_list:
                document.add_paragraph(value, style="List Bullet")
        else:
            document.add_paragraph(empty_text)

    docx_list(
        "Équipements de protection individuelle (EPI)",
        _list_names(version, "version_epis", "epi"),
        "Aucun EPI renseigné.",
    )
    docx_list(
        "Équipements de protection collective (EPC)",
        _list_names(version, "version_epcs", "epc"),
        "Aucun EPC renseigné.",
    )
    docx_list(
        "Risques",
        _list_names(version, "version_risques", "risque"),
        "Aucun risque renseigné.",
    )
    docx_list(
        "Outillages",
        _list_quantified(version, "version_outillages", "outillage"),
        "Aucun outillage renseigné.",
    )
    docx_list(
        "Pièces de rechange",
        _list_quantified(version, "version_pieces_rechange", "piece"),
        "Aucune pièce de rechange renseignée.",
    )

    document.add_heading("Étapes & actions", level=1)
    etapes = _ordered_manager_items(version, "etapes", "ordre", "numero")
    if not etapes:
        document.add_paragraph("Aucune étape renseignée.")
    else:
        for etape in etapes:
            document.add_heading(
                f"{getattr(etape, 'numero', '-')}. {_text(getattr(etape, 'titre', None))}",
                level=2,
            )
            document.add_paragraph(
                f"Durée : {getattr(etape, 'duree_minutes', 0) or 0} minutes"
            )
            if getattr(etape, "description", None):
                document.add_paragraph(_text(etape.description))

            for action in _ordered_manager_items(etape, "actions", "ordre"):
                document.add_paragraph(
                    _text(getattr(action, "contenu", None)),
                    style="List Bullet",
                )

            for image in _ordered_manager_items(etape, "images", "ordre"):
                raw_image = _image_bytes_from_value(getattr(image, "image_url", None))
                if raw_image:
                    try:
                        document.add_picture(BytesIO(raw_image), width=Inches(5.7))
                    except Exception:
                        pass
                if getattr(image, "description", None):
                    document.add_paragraph(_text(image.description))

    document.add_heading("Recommandations", level=1)
    recommendations = _manager_items(version, "recommandations")
    if recommendations:
        for recommendation in recommendations:
            if getattr(recommendation, "titre", None):
                document.add_heading(_text(recommendation.titre), level=2)
            document.add_paragraph(_text(getattr(recommendation, "contenu", None), ""))
    else:
        document.add_paragraph("Aucune recommandation renseignée.")

    document.add_heading("Documents liés", level=1)
    linked_documents = _manager_items(version, "documents")
    if linked_documents:
        for linked_document in linked_documents:
            text = _text(getattr(linked_document, "titre", None))
            if getattr(linked_document, "reference", None):
                text += f" — Réf. {linked_document.reference}"
            document.add_paragraph(text, style="List Bullet")
            if getattr(linked_document, "description", None):
                document.add_paragraph(_text(linked_document.description))
            if getattr(linked_document, "fichier_url", None):
                document.add_paragraph(f"Fichier : {_text(linked_document.fichier_url)}")
    else:
        document.add_paragraph("Aucun document lié.")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue(), _filename(version, "docx")