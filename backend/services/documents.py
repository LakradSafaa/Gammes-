import base64
from io import BytesIO

import qrcode

from django.conf import settings
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.template.loader import render_to_string

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from gammes_maintenance.models import FichierGenere


# ============================================================
# WEASYPRINT
# ============================================================

try:
    from weasyprint import HTML
except (ImportError, OSError):
    HTML = None


# ============================================================
# LIBELLES
# ============================================================

STATUT_LABELS = {
    "brouillon": "Brouillon",
    "en_validation": "En validation",
    "validee": "Validée",
    "archivee": "Archivée",
}


TYPE_MAINTENANCE_LABELS = {
    "preventif": "Préventive",
    "correctif": "Corrective",
    "amelioratif": "Améliorative",
}


GREEN_DARK = "14532D"
GREEN_MAIN = "16A34A"


# ============================================================
# FONCTIONS UTILITAIRES
# ============================================================

def get_statut_label(version):
    """
    Retourne le libellé du statut.
    """

    statut = getattr(version, "statut", None)

    if not statut:
        return "-"

    return STATUT_LABELS.get(
        statut,
        str(statut).replace("_", " ").capitalize(),
    )


def get_type_maintenance_label(version):
    """
    Retourne le libellé du type de maintenance.
    """

    type_maintenance = getattr(
        version,
        "type_maintenance",
        None,
    )

    if not type_maintenance:
        return "-"

    return TYPE_MAINTENANCE_LABELS.get(
        type_maintenance,
        str(type_maintenance)
        .replace("_", " ")
        .capitalize(),
    )


def get_qr_url(version):
    """
    Retourne l'URL publique dynamique de la gamme.

    Le QR pointe vers le code de la gamme et non vers une version
    particulière. Il reste donc valable lorsque V1 devient V2, V3, etc.
    """

    app_base_url = getattr(
        settings,
        "APP_BASE_URL",
        "http://localhost:5173",
    ).rstrip("/")

    return (
        f"{app_base_url}/qr/"
        f"{version.gamme.code}"
    )


def save_generated_file(
    version,
    file_type,
    filename,
    data,
):
    """
    Enregistre le fichier généré avec le storage Django
    puis crée l'enregistrement dans fichiers_generes.
    """

    storage_path = (
        f"generated/"
        f"{version.gamme.code}/"
        f"{version.code_version}/"
        f"{filename}"
    )

    if default_storage.exists(storage_path):
        default_storage.delete(storage_path)

    saved_path = default_storage.save(
        storage_path,
        ContentFile(data),
    )

    try:
        file_url = default_storage.url(saved_path)
    except Exception:
        file_url = saved_path

    fichier = FichierGenere.objects.create(
        version=version,
        type_fichier=file_type,
        fichier_url=file_url,
        nom_fichier=filename,
    )

    return fichier


def set_cell_border_none(cell):
    """
    Supprime les bordures d'une cellule Word.
    """

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_borders = tc_pr.first_child_found_in("w:tcBorders")

    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)

        element.set(qn("w:val"), "nil")


def configure_word_styles(document):
    """
    Applique une identité visuelle verte et blanche au document Word.
    """

    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(31, 41, 55)

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(20, 83, 45)

    for style_name, size in (
        ("Heading 1", 14),
        ("Heading 2", 12),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(22, 163, 74)


# ============================================================
# QR CODE
# ============================================================

def qr_png_bytes(version):
    """
    Génère le QR Code dynamique de la gamme.
    """

    url = get_qr_url(version)

    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=4,
    )

    qr.add_data(url)
    qr.make(fit=True)

    image = qr.make_image(
        fill_color="black",
        back_color="white",
    )

    buffer = BytesIO()

    image.save(
        buffer,
        format="PNG",
    )

    buffer.seek(0)

    return buffer.getvalue()


# ============================================================
# PDF
# ============================================================

def generate_pdf(version):
    """
    Génère le PDF d'une version.
    Le QR Code est affiché dans le coin supérieur droit
    de la première page.
    """

    if HTML is None:
        raise RuntimeError(
            "WeasyPrint n'est pas disponible sur cette machine. "
            "Sous Windows, les dépendances système GTK/Pango "
            "doivent être installées avant la génération PDF."
        )

    qr_data_url = None
    qr_url = get_qr_url(version)

    try:
        qr_data = qr_png_bytes(version)
        qr_base64 = base64.b64encode(qr_data).decode("ascii")
        qr_data_url = (
            "data:image/png;base64,"
            f"{qr_base64}"
        )
    except Exception:
        qr_data_url = None

    html = render_to_string(
        "gammes/version_pdf.html",
        {
            "version": version,
            "statut_label": get_statut_label(version),
            "type_maintenance_label":
                get_type_maintenance_label(version),
            "qr_data_url": qr_data_url,
            "qr_url": qr_url,
        },
    )

    pdf_data = HTML(
        string=html,
        base_url=str(settings.BASE_DIR),
    ).write_pdf()

    filename = (
        f"{version.gamme.code}_"
        f"{version.code_version}.pdf"
    )

    return save_generated_file(
        version=version,
        file_type="pdf",
        filename=filename,
        data=pdf_data,
    )


# ============================================================
# WORD
# ============================================================

def generate_docx(version):
    """
    Génère le document Word d'une version de gamme opératoire.
    Le QR Code est positionné en haut à droite de la première page.
    """

    document = Document()
    configure_word_styles(document)

    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    # ========================================================
    # EN-TÊTE PREMIÈRE PAGE AVEC QR CODE
    # ========================================================

    header_table = document.add_table(
        rows=1,
        cols=2,
    )

    header_table.autofit = False

    left_cell = header_table.cell(0, 0)
    right_cell = header_table.cell(0, 1)

    left_cell.width = Inches(5.65)
    right_cell.width = Inches(1.35)

    set_cell_border_none(left_cell)
    set_cell_border_none(right_cell)

    title_paragraph = left_cell.paragraphs[0]
    title_run = title_paragraph.add_run(
        f"Gamme opératoire - {version.gamme.code}"
    )
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(20, 83, 45)

    designation_paragraph = left_cell.add_paragraph()
    designation_run = designation_paragraph.add_run(
        f"{version.gamme.designation or '-'}"
    )
    designation_run.font.name = "Arial"
    designation_run.font.size = Pt(11)

    meta_paragraph = left_cell.add_paragraph()
    meta_run = meta_paragraph.add_run(
        f"Version {version.code_version} | "
        f"{get_statut_label(version)}"
    )
    meta_run.bold = True
    meta_run.font.name = "Arial"
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = RGBColor(22, 163, 74)

    if getattr(version, "date_version", None):
        date_paragraph = left_cell.add_paragraph(
            f"Date de version : {version.date_version}"
        )
        date_paragraph.runs[0].font.size = Pt(9)

    qr_paragraph = right_cell.paragraphs[0]
    qr_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    try:
        qr_data = qr_png_bytes(version)
        qr_buffer = BytesIO(qr_data)

        qr_run = qr_paragraph.add_run()
        qr_run.add_picture(
            qr_buffer,
            width=Inches(1.15),
        )

        qr_label = right_cell.add_paragraph()
        qr_label.alignment = WD_ALIGN_PARAGRAPH.CENTER

        qr_label_run = qr_label.add_run(
            "Scanner la gamme"
        )
        qr_label_run.bold = True
        qr_label_run.font.name = "Arial"
        qr_label_run.font.size = Pt(8)
        qr_label_run.font.color.rgb = RGBColor(20, 83, 45)

    except Exception:
        qr_error = right_cell.add_paragraph(
            "QR indisponible"
        )
        qr_error.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph()

    # ========================================================
    # INFORMATIONS GÉNÉRALES
    # ========================================================

    document.add_heading(
        "Informations générales",
        level=1,
    )

    document.add_paragraph(
        "Type de maintenance : "
        f"{get_type_maintenance_label(version)}"
    )

    document.add_paragraph(
        f"Périodicité : "
        f"{version.periodicite or '-'}"
    )

    document.add_paragraph(
        f"Main-d'œuvre : "
        f"{version.main_oeuvre or 0}"
    )

    document.add_paragraph(
        f"Durée totale : "
        f"{version.duree_minutes or 0} minutes"
    )

    document.add_paragraph(
        f"Rédacteur : "
        f"{version.redacteur or '-'}"
    )

    document.add_paragraph(
        f"Valideur : "
        f"{version.valideur or '-'}"
    )

    if version.modifications:
        document.add_paragraph(
            f"Modifications : "
            f"{version.modifications}"
        )

    # ========================================================
    # PARAMÈTRES
    # ========================================================

    document.add_heading(
        "Paramètres d'intervention",
        level=1,
    )

    document.add_paragraph(
        f"Référentiel : "
        f"{'Oui' if version.referentiel else 'Non'}"
    )

    document.add_paragraph(
        f"Rapport : "
        f"{'Oui' if version.rapport else 'Non'}"
    )

    document.add_paragraph(
        f"Production : "
        f"{'Oui' if version.production else 'Non'}"
    )

    document.add_paragraph(
        f"Arrêt : "
        f"{'Oui' if version.arret else 'Non'}"
    )

    document.add_paragraph(
        f"Mode dégradé : "
        f"{'Oui' if version.degrade else 'Non'}"
    )

    # ========================================================
    # RISQUES
    # ========================================================

    document.add_heading(
        "Risques",
        level=1,
    )

    risques = (
        version.version_risques
        .select_related("risque")
        .all()
    )

    if risques.exists():

        for association in risques:

            document.add_paragraph(
                association.risque.nom,
                style="List Bullet",
            )

    else:

        document.add_paragraph(
            "Aucun risque renseigné."
        )

    # ========================================================
    # EPI
    # ========================================================

    document.add_heading(
        "Équipements de protection individuelle",
        level=1,
    )

    epis = (
        version.version_epis
        .select_related("epi")
        .all()
    )

    if epis.exists():

        for association in epis:

            document.add_paragraph(
                association.epi.nom,
                style="List Bullet",
            )

    else:

        document.add_paragraph(
            "Aucun EPI renseigné."
        )

    # ========================================================
    # OUTILLAGES
    # ========================================================

    document.add_heading(
        "Outillages",
        level=1,
    )

    outillages = (
        version.version_outillages
        .select_related("outillage")
        .all()
    )

    if outillages.exists():

        for association in outillages:

            quantite = (
                association.quantite
                if association.quantite is not None
                else 1
            )

            document.add_paragraph(
                f"{association.outillage.nom} "
                f"× {quantite}",
                style="List Bullet",
            )

    else:

        document.add_paragraph(
            "Aucun outillage renseigné."
        )

    # ========================================================
    # PIÈCES DE RECHANGE
    # ========================================================

    document.add_heading(
        "Pièces de rechange",
        level=1,
    )

    pieces = (
        version.version_pieces_rechange
        .select_related("piece")
        .all()
    )

    if pieces.exists():

        for association in pieces:

            quantite = (
                association.quantite
                if association.quantite is not None
                else 1
            )

            code_piece = (
                association.piece.code or ""
            )

            if code_piece:

                texte = (
                    f"{code_piece} - "
                    f"{association.piece.nom} "
                    f"× {quantite}"
                )

            else:

                texte = (
                    f"{association.piece.nom} "
                    f"× {quantite}"
                )

            document.add_paragraph(
                texte,
                style="List Bullet",
            )

    else:

        document.add_paragraph(
            "Aucune pièce de rechange renseignée."
        )

    # ========================================================
    # ÉTAPES ET ACTIONS
    # ========================================================

    document.add_heading(
        "Étapes et actions",
        level=1,
    )

    etapes = (
        version.etapes
        .prefetch_related(
            "actions",
            "images",
        )
        .order_by(
            "ordre",
            "numero",
        )
    )

    if etapes.exists():

        for etape in etapes:

            document.add_heading(
                f"{etape.numero}. "
                f"{etape.titre}",
                level=2,
            )

            document.add_paragraph(
                f"Durée : "
                f"{etape.duree_minutes or 0} minutes"
            )

            if etape.description:

                document.add_paragraph(
                    etape.description
                )

            actions = (
                etape.actions
                .all()
                .order_by("ordre")
            )

            if actions.exists():

                document.add_paragraph(
                    "Actions :"
                )

                for action in actions:

                    document.add_paragraph(
                        f"{action.ordre}. "
                        f"{action.contenu}",
                        style="List Bullet",
                    )

            else:

                document.add_paragraph(
                    "Aucune action renseignée."
                )

            images = (
                etape.images
                .all()
                .order_by("ordre")
            )

            if images.exists():

                document.add_paragraph(
                    "Images associées :"
                )

                for image in images:

                    texte_image = (
                        image.description
                        or image.image_url
                        or "Image"
                    )

                    document.add_paragraph(
                        texte_image,
                        style="List Bullet",
                    )

    else:

        document.add_paragraph(
            "Aucune étape renseignée."
        )

    # ========================================================
    # RECOMMANDATIONS
    # ========================================================

    document.add_heading(
        "Recommandations",
        level=1,
    )

    recommandations = (
        version.recommandations.all()
    )

    if recommandations.exists():

        for recommandation in recommandations:

            if recommandation.titre:

                document.add_heading(
                    recommandation.titre,
                    level=2,
                )

            document.add_paragraph(
                recommandation.contenu
            )

    else:

        document.add_paragraph(
            "Aucune recommandation renseignée."
        )

    # ========================================================
    # DOCUMENTS LIÉS
    # ========================================================

    document.add_heading(
        "Documents liés",
        level=1,
    )

    documents_lies = (
        version.documents.all()
    )

    if documents_lies.exists():

        for document_lie in documents_lies:

            texte = document_lie.titre

            if document_lie.reference:

                texte += (
                    f" - Référence : "
                    f"{document_lie.reference}"
                )

            document.add_paragraph(
                texte,
                style="List Bullet",
            )

            if document_lie.description:

                document.add_paragraph(
                    document_lie.description
                )

            if document_lie.fichier_url:

                document.add_paragraph(
                    f"Fichier : "
                    f"{document_lie.fichier_url}"
                )

    else:

        document.add_paragraph(
            "Aucun document lié."
        )

    # ========================================================
    # DURÉE TOTALE
    # ========================================================

    document.add_heading(
        "Durée totale",
        level=1,
    )

    document.add_paragraph(
        f"{version.duree_minutes or 0} minutes"
    )

    # ========================================================
    # LIEN QR
    # ========================================================

    qr_info = document.add_paragraph()
    qr_info.alignment = WD_ALIGN_PARAGRAPH.CENTER

    qr_info_run = qr_info.add_run(
        f"QR dynamique : {get_qr_url(version)}"
    )
    qr_info_run.font.size = Pt(8)
    qr_info_run.font.color.rgb = RGBColor(100, 116, 139)

    # ========================================================
    # CRÉATION DU FICHIER WORD
    # ========================================================

    buffer = BytesIO()

    document.save(buffer)

    buffer.seek(0)

    filename = (
        f"{version.gamme.code}_"
        f"{version.code_version}.docx"
    )

    return save_generated_file(
        version=version,
        file_type="word",
        filename=filename,
        data=buffer.getvalue(),
    )
